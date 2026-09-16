from pathlib import Path
import chardet
import pdfplumber
from docx import Document as DocxDocument


class DocumentParser:
    def parse(self, file_path: str) -> str:
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf(file_path)
        elif ext == ".docx":
            return self._parse_docx(file_path)
        elif ext in (".txt", ".md"):
            return self._parse_txt(file_path)
        elif ext == ".doc":
            raise ValueError("旧版 .doc 不受支持，请另存为 .docx")
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _parse_pdf(self, file_path: str) -> str:
        parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                # 1. 提取页面文本
                text = page.extract_text()
                if text: #只保留非空文本（if text:）
                    parts.append(text)
                # 2. 提取表格
                tables = page.extract_tables()
                for table in tables:
                    if table:
                        table_text = self._format_table(table)
                        parts.append(table_text)#表格作为独立片段添加到文本中
        return "\n\n".join(parts)

    def _parse_docx(self, file_path: str) -> str:
        doc = DocxDocument(file_path)
        parts = []
        for para in doc.paragraphs: #遍历所有段落
            if para.text.strip(): #过滤掉空白段落（去除首尾空格后为空）
                parts.append(para.text)
        for table in doc.tables:#遍历所有表格
            rows = []
            for row in table.rows:#遍历表格的行
                cells = [cell.text.strip() for cell in row.cells]#获取每个单元格的文本
                rows.append(cells)
            if rows:
                parts.append(self._format_table(rows))
        return "\n\n".join(parts)

    def _parse_txt(self, file_path: str) -> str:
        with open(file_path, "rb") as f:#以 rb 模式读取原始字节
            raw = f.read()
        detected = chardet.detect(raw)#使用 chardet 库自动检测编码
        encoding = detected["encoding"] or "utf-8"#降级策略: 如果检测失败，默认使用 UTF-8
        return raw.decode(encoding)

    #统一的表格转文本方法，将二维数组转换为可读的文本格式
    def _format_table(self, table: list[list[str]]) -> str:
        lines = []
        for row in table:
            row = [cell if cell else "" for cell in row]#处理None
            lines.append(" | ".join(row))# 用 | 分隔列
        return "\n".join(lines)# 用换行分隔行
